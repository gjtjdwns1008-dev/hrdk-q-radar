# -*- coding: utf-8 -*-
"""
[2/3] run.py — 재분석 실행  (시트 읽기만, 쓰기 없음)
=====================================================
대상목록(scan 산출) 의 재분석대상=Y 행을 시행일자별로 묶어
  ① 그 시행일자 판본의 원문(별표 포함)을 법제처에서 다시 받고     … core.scraper.get_base_laws(only_names=)
  ② 개정 프롬프트로 분석하고                                        … pipeline.brain.run_ai_analysis
  ③ 종목 사전 필터·워크넷 스위치를 일일 파이프라인과 똑같이 적용해  … core.certs / core.worknet
  ④ 구(대장 현재값) / 신(재분석값) 을 나란히 놓은 대조 엑셀을 만든다.  → work/재분석_대조_YYYYMMDD.xlsx

특징
  · 캐시: work/cache/laws_{시행일자}.json (원문), work/cache/result_{MST_ID}.json (분석 결과)
    → 중단 후 다시 실행하면 이미 된 것은 건너뛴다(법제처·AI 재호출 없음). 엑셀은 캐시 전체로 매번 새로 만든다.
  · 실패는 감추지 않고 '실패' 시트에 단계·사유를 적는다(원문 없음 / 판본 목록에 없음 / AI 실패).
  · --mock : 법제처·AI 를 부르지 않고 가짜 값으로 엑셀 생성 경로만 검증하는 테스트 모드(산출 파일명에 _MOCK).

사용
  python run.py --targets work/대상목록_20260916.xlsx            # 전체
  python run.py --targets ... --limit 20                          # 앞 20건만(비용·시간 실측)
  python run.py --targets ... --date 20260102                     # 특정 시행일자만
  python run.py --targets ... --local 대장.xlsx                   # 구 값을 로컬 xlsx 에서
"""
from __future__ import annotations

import argparse
import json, traceback
import os
import sys
import time
from collections import Counter, defaultdict
from pathlib import Path

HERE = Path(__file__).resolve().parent
REPO = HERE.parent.parent
sys.path.insert(0, str(HERE))
sys.path.insert(0, str(REPO / "pipeline"))       # brain
sys.path.insert(0, str(REPO / "core"))           # hrdk_law_core (pip 설치돼 있으면 그쪽이 우선)
import common as C  # noqa: E402

CACHE = C.WORK / "cache"
CACHE.mkdir(exist_ok=True)


def preflight() -> bool:
    """실행 전 점검: 패치된 core·brain 을 불러오는지 확인. 하나라도 옛 파일이면 어디를 덮어쓸지 알려주고 멈춘다."""
    import inspect
    ok = True
    print("── 사전 점검 ──")
    try:
        import hrdk_law_core, hrdk_law_core.scraper as sc, hrdk_law_core.llm_client as lc, hrdk_law_core.certs as ce
        core_dir = Path(hrdk_law_core.__file__).parent
        print(f"· core 경로  : {core_dir}")
        if "only_names" not in inspect.signature(sc.get_base_laws).parameters:
            print(f"  ✗ scraper.py 가 패치 전 파일입니다 → {core_dir / 'scraper.py'} 를 패치본으로 덮어쓰세요."); ok = False
        else:
            print("  ✓ scraper.py (only_names 지원)")
        if "json_mode" not in inspect.signature(lc.LLMClient.generate).parameters:
            print(f"  ✗ llm_client.py 가 패치 전 파일입니다 → {core_dir / 'llm_client.py'} 를 덮어쓰세요."); ok = False
        else:
            print("  ✓ llm_client.py (json_mode·thinking_level 지원)")
        if not hasattr(ce, "SCOPE_TOKEN_ALL"):
            print(f"  ✗ certs.py 가 패치 전 파일입니다 → {core_dir / 'certs.py'} 를 덮어쓰세요."); ok = False
        else:
            print("  ✓ certs.py ([전종목] 토큰 지원)")
        if str(core_dir).replace("\\", "/").find("/site-packages/") >= 0:
            print("  ⚠ pip 에 설치된 core 를 불러왔습니다. 레포 폴더의 core/hrdk_law_core/ 가 있는지 확인하세요"
                  f" (기대 위치: {REPO / 'core' / 'hrdk_law_core'}).")
    except Exception as e:
        print(f"  ✗ hrdk_law_core 를 불러올 수 없습니다: {e}"); ok = False
    try:
        import brain
        print(f"· brain 경로 : {brain.__file__}")
        if not hasattr(brain, "_apply_review_guards"):
            print(f"  ✗ brain.py 가 패치 전 파일입니다 → {brain.__file__} 를 덮어쓰세요."); ok = False
        else:
            print("  ✓ brain.py (개정 프롬프트·후처리 방어망)")
    except Exception as e:
        print(f"  ✗ pipeline/brain.py 를 불러올 수 없습니다: {e}"); ok = False
    m, lv = os.environ.get("LLM_MODEL", ""), os.environ.get("LLM_THINKING_LEVEL", "")
    if not m or not lv:
        print(f"  ⚠ local_run/local.env 에 {'LLM_MODEL ' if not m else ''}{'LLM_THINKING_LEVEL ' if not lv else ''}이(가) 없습니다."
              " 기본값(3.5-flash·thinking 미지정)으로 돌아갑니다. 3.8 로 돌리려면 local_run/local.env 에 두 줄을 추가하세요:\n"
              "     LLM_MODEL=gemini-3.8-flash\n     LLM_THINKING_LEVEL=medium")
        if len(sys.argv) == 1:
            ans = input("  그래도 기본 모델로 계속할까요? (y = 계속, 그 외 = 중단): ").strip().lower()
            if ans != "y":
                ok = False
    print("──────────────")
    return ok


# ─────────────────────────────────────────────────────────────
# 1. 원문 (시행일자 판본) — 캐시 우선
# ─────────────────────────────────────────────────────────────
def fetch_laws_for_date(date: str, names: set[str], mock: bool) -> list[dict] | None:
    """해당 시행일자에 시행된 법령 중 names(정규화 법령명) 만 원문 포함으로. None = 법제처 연결 실패."""
    cp = CACHE / f"laws_{date}.json"
    if cp.exists():
        laws = json.loads(cp.read_text(encoding="utf-8"))
        have = {C.norm_law_name(l["법령명"]) for l in laws}
        if names <= have:                          # 이번 대상이 캐시에 다 있으면 재수집 없음
            return laws
        names = names - have                       # 새로 필요한 것만 받아 캐시에 합침
        base = laws
    else:
        base = []
    if mock:
        laws = [{"법령명": n, "시행일자": date, "소관부처": "(mock)", "공포번호": "", "공포일자": "",
                 "원본": f"(MOCK 원문) {n} 제1조 국가기술자격 취득자를 우대한다.", "링크": f"https://www.law.go.kr/법령/{n}"}
                for n in names]
    else:
        from hrdk_law_core.scraper import get_base_laws
        laws = get_base_laws(api_key=os.environ["LAW_API_KEY"], target_date=date, only_names=names)
        if laws is None:
            return None
    merged = base + [l for l in laws if C.norm_law_name(l["법령명"]) not in {C.norm_law_name(b["법령명"]) for b in base}]
    cp.write_text(json.dumps(merged, ensure_ascii=False), encoding="utf-8")
    return merged


ERROR_LOG = C.WORK / f"오류로그_{C.STAMP}.txt"


def prompt_version() -> str:
    """★2026-10-14 brain.PROMPT_VERSION — 없는 구판 brain 이면 '(미표시)'."""
    try:
        import brain
        return str(getattr(brain, "PROMPT_VERSION", "(미표시)"))
    except Exception:
        return "(미표시)"


def cache_fresh(mid: str, pver: str) -> bool:
    """★2026-10-14 결과 캐시가 있고, 그때의 프롬프트 버전이 지금과 같을 때만 재사용."""
    cp = CACHE / f"result_{mid}.json"
    if not cp.exists():
        return False
    try:
        return json.loads(cp.read_text(encoding="utf-8")).get("_프롬프트버전", "(미표시)") == pver
    except Exception:
        return False


def unique_path(path: Path) -> Path:
    """★2026-10-14 같은 이름이 이미 있거나 열려 있으면 _2, _3 … 을 붙여 새 파일로 저장한다(덮어쓰기·잠금 방지)."""
    cand, n = path, 1
    while True:
        if not cand.exists():
            return cand
        try:
            with open(cand, "r+b"):          # 열 수 있으면(엑셀이 안 잡고 있으면) 그래도 덮어쓰지 않고 번호를 붙인다
                pass
        except PermissionError:
            pass
        n += 1
        cand = path.with_name(f"{path.stem}_{n}{path.suffix}")


def log_exception(mid: str, name: str, date: str, stage: str, law: dict | None, exc: Exception) -> str:
    """★2026-10-14 법령 단위 격리: 한 건의 예외를 실행 전체의 죽음으로 키우지 않는다.
    전체 traceback 과 법령 정보를 work/오류로그_날짜.txt 에 덧붙이고, 실패 사유용 한 줄을 돌려준다."""
    why = f"{stage} 예외 {type(exc).__name__}: {str(exc)[:120]}"
    try:
        with open(ERROR_LOG, "a", encoding="utf-8") as f:
            f.write(f"\n{'=' * 70}\n{time.strftime('%Y-%m-%d %H:%M:%S')}  {mid}  {name}  시행일자 {date}  단계 {stage}\n")
            if law:
                f.write(f"판본출처: {law.get('판본출처', '(그날 현행 목록)')} · 원문 {len(str(law.get('원본', ''))):,}자 · 스킵여부 {law.get('스킵여부')}\n")
            f.write(traceback.format_exc())
    except Exception:
        pass
    return why


def fetch_version_fallback(name: str, date: str, day_laws: list) -> dict | str | None:
    """★2026-10-14 현행 목록에 없는 법령을 연혁(eflaw)에서 '시행일자 판본'으로 받는다. 받은 판본은 그 날짜 캐시
    (laws_{date}.json) 에 합쳐 두어 다시 실행해도 법제처를 다시 부르지 않는다.
    반환: dict(원문 레코드, '판본출처' 포함) / str(사유: 판본 없음 등) / None(연결 실패)"""
    from hrdk_law_core.scraper import get_law_version
    rec = get_law_version(os.environ["LAW_API_KEY"], name, date)
    if isinstance(rec, dict):
        rec.setdefault("판본출처", "eflaw 연혁 판본")
        day_laws.append(rec)                                   # 호출 측 by_name 과 같은 목록 객체에 추가
        cp = CACHE / f"laws_{date}.json"
        cp.write_text(json.dumps(day_laws, ensure_ascii=False), encoding="utf-8")
    return rec


# ─────────────────────────────────────────────────────────────
# 2. 분석 — 일일 파이프라인(main.py 160~190행)과 같은 순서
# ─────────────────────────────────────────────────────────────
def analyze_one(law: dict, mock_old: dict | None) -> tuple[bool, dict, str]:
    """(성공여부, 신 값 dict[대장 열명→값], 실패사유)"""
    import brain
    from hrdk_law_core.certs import get_relevant_certs_text, normalize_cert_string
    from hrdk_law_core.worknet import get_worknet_job_count

    year = int(C.norm_date(law.get("시행일자", ""))[:4]) if C.norm_date(law.get("시행일자", "")) else None
    if mock_old is not None:
        _install_mock_llm(mock_old)
    ok, is_related, info = brain.run_ai_analysis(law, get_relevant_certs_text(law.get("원본", ""), year=year))
    if not ok:
        return False, {}, "AI 분석 실패(재시도 소진 또는 JSON 파싱 불가)"
    std, dropped = normalize_cert_string(info.get("관련 종목", ""), year=year)
    info["관련 종목"] = std
    if dropped:
        info["_사전밖_제외"] = ", ".join(dropped)          # 대조표 참고용(시트에 쓰지 않음)
    info["워크넷 실시간 구인건수"] = get_worknet_job_count(std, api_key=os.environ.get("WORKNET_API_KEY", "")) \
        if info.get("우대여부") == "O" else "-"
    return True, {c: str(info.get(c, "")) for c in C.ANALYSIS_COLS} | {"_사전밖_제외": info.get("_사전밖_제외", "")}, ""


def _install_mock_llm(old: dict):
    """★MOCK: 대장 현재값을 거의 그대로 돌려주되 두 칸만 바꿔 '변경 표시'가 보이게 한다. 실제 분석 아님."""
    import brain
    canned = {
        "연관도": old.get("연관도") or "단순관련", "개정유형": old.get("개정유형") or "일부개정",
        "주요_제개정내용": old.get("주요 제·개정내용") or "- (mock)", "종목": old.get("관련 종목", ""),
        "활용도_구분": old.get("활용도_구분", ""), "활용도_상세": old.get("활용도_상세", ""),
        "우대여부": old.get("우대여부") or "X", "우대분류": old.get("우대분류") or "기타",
        "Track1_취급유형": (old.get("Track1_취급유형") or "Z")[:1], "Track1_위험도": (old.get("Track1_위험도") or "X")[:1],
        "Track2_효용코드": (old.get("Track2_효용코드") or "Ⅳ-0").split(" ")[0], "중처법대상": old.get("중처법대상") or "비대상",
        "조문_요약": old.get("조문 요약", ""), "상세_분석": old.get("상세 분석 결과", ""),
        "AI_신뢰도": old.get("AI신뢰도") or "높음", "검토필요": "O", "검토사유": "(MOCK) 테스트용 변경 표시", "조문리스트": [],
    }
    if canned["우대여부"] == "O" and canned["우대분류"] == "기타":
        canned["우대분류"] = "인사우대"              # A 유형 행은 재분류된 것처럼 보이게(테스트용)
    class _F:
        def generate_with_retry(self, prompt, **kw): return json.dumps(canned, ensure_ascii=False)
    brain._llm = _F()


# ─────────────────────────────────────────────────────────────
# 3. 메인
# ─────────────────────────────────────────────────────────────
def main():
    ap = argparse.ArgumentParser(description="재분석 실행 (시트 쓰기 없음)")
    ap.add_argument("--targets", help="scan 산출 대상목록 xlsx (기본: work/ 의 최신)")
    ap.add_argument("--local", help="구 값(대장 현재값)을 읽을 로컬 대장 xlsx")
    ap.add_argument("--limit", type=int, default=0, help="앞 N건만")
    ap.add_argument("--date", help="특정 시행일자(YYYYMMDD)만")
    ap.add_argument("--mock", action="store_true", help="★테스트 모드: 법제처·AI 호출 없음")
    ap.add_argument("--force", action="store_true", help="★2026-10-14 캐시된 결과가 있어도 다시 분석(반복 일치율 측정용)")
    args = ap.parse_args()
    if len(sys.argv) == 1:                                   # ▶ 버튼(인자 없음): 화면에서 묻기
        print("── 재분석 실행 (시트에는 쓰지 않음) ──")
        ans = input("몇 건만 돌릴까요? 숫자 입력, 그냥 엔터 = 전량: ").strip()
        args.limit = int(ans) if ans.isdigit() else 0
        ans = input("이미 분석된(캐시) 건도 다시 분석할까요? y = 다시 분석, 엔터 = 캐시 재사용: ").strip().lower()
        args.force = ans in ("y", "yes", "ㅛ")

    C.load_env()
    if args.local:
        os.environ["LOCAL_XLSX"] = args.local
    args.mock = args.mock or os.environ.get("QRADAR_TOOL_MOCK") == "1"      # 테스트 전용 훅
    need = [] if os.environ.get("LOCAL_XLSX") else ["GOOGLE_SHEET_URL", "GCP_SERVICE_ACCOUNT_JSON"]
    if not args.mock:
        need += ["LAW_API_KEY", "GEMINI_API_KEY"]
    if not C.require(need, "재분석 실행"):
        sys.exit(1)
    os.environ.setdefault("WORKNET_ENABLED", "0")                 # 결정('26.9.16): 실시간 조회 배제
    if not preflight():
        print("⛔ 사전 점검 실패 — 위 안내대로 파일·local_run/local.env 를 고친 뒤 다시 실행하세요."); sys.exit(1)

    tp = Path(args.targets) if args.targets else sorted(C.WORK.glob("대상목록_*.xlsx"))[-1]
    import pandas as pd
    tg = pd.read_excel(tp, sheet_name="대상목록", dtype=str).fillna("")
    tg = tg[tg["재분석대상"] == "Y"]
    if args.date:
        tg = tg[tg["시행일자"].map(C.norm_date) == args.date]
    if args.limit:
        tg = tg.head(args.limit)
    print(f"· 대상: {tp.name} → {len(tg)}건" + (" ★MOCK" if args.mock else ""))
    if not len(tg):
        sys.exit(0)

    header, rows = C.read_ledger()
    old_by_id = {r["MST_ID"]: r for r in rows if r.get("MST_ID")}

    model, lvl = os.environ.get("LLM_MODEL", "(기본 gemini-3.5-flash)"), os.environ.get("LLM_THINKING_LEVEL", "(미지정)")
    pver = prompt_version()
    print(f"· 모델 {model} · thinking_level {lvl} · 프롬프트 {pver} · WORKNET_ENABLED {os.environ.get('WORKNET_ENABLED')}")
    t0 = time.time()
    failures, done, skipped = [], 0, 0
    by_date = defaultdict(list)
    for _, t in tg.iterrows():
        by_date[C.norm_date(t["시행일자"])].append(t)

    for di, (date, items) in enumerate(sorted(by_date.items()), 1):
        pending = [t for t in items if args.force or not cache_fresh(t["MST_ID"], pver)]   # ★2026-10-14 버전 다르면 다시
        skipped += len(items) - len(pending)
        if not pending:
            continue
        names = {C.norm_law_name(t["법령명"]) for t in pending}
        print(f"\n[{di}/{len(by_date)}] 시행일자 {date} — {len(pending)}건 (캐시 제외)")
        laws = fetch_laws_for_date(date, names, args.mock)
        if laws is None:
            for t in pending:
                failures.append((t["MST_ID"], t["법령명"], date, "원문", "법제처 연결 실패(재시도 소진) — 다시 실행하면 재시도"))
            print("  ❌ 법제처 연결 실패 — 이 날짜는 통째로 실패 기록"); continue
        by_name = {C.norm_law_name(l["법령명"]): l for l in laws}
        for t in pending:
            mid, name = t["MST_ID"], t["법령명"]
            law = by_name.get(C.norm_law_name(name))
            if law is None and not args.mock:
                # ★2026-10-14 연혁 판본 보강: 그날 시행 목록(target=law)은 '현행' 법령만 주므로, 그 뒤 다시 개정된
                #   법령의 옛 판본은 목록에 없다 → 법령명으로 연혁(eflaw)을 검색해 그 시행일자 판본을 받는다.
                try:
                    law = fetch_version_fallback(name, date, laws)
                except Exception as e:                   # ★2026-10-14 격리
                    failures.append((mid, name, date, "원문", log_exception(mid, name, date, "연혁 판본 수집", None, e)))
                    print(f"  ❌ {mid} 연혁 판본 수집 예외 → 오류로그 기록, 다음 법령 진행"); continue
                if isinstance(law, str):                 # 사유 문자열 = 연혁에서도 못 찾음
                    failures.append((mid, name, date, "원문", f"현행 목록에 없음 → 연혁 검색: {law}")); continue
                if law is None:
                    failures.append((mid, name, date, "원문", "현행 목록에 없음 → 연혁 검색 중 법제처 연결 실패 — 다시 실행하면 재시도")); continue
            if law is None:
                failures.append((mid, name, date, "원문", "해당 시행일자 판본 목록에 법령명 없음(명칭 변경·폐지·시행일자 오기 의심)")); continue
            if law.get("스킵여부"):
                failures.append((mid, name, date, "원문", f"수집 단계 스킵 대상({law.get('스킵사유','직제 등')}) — 재분석 불가")); continue
            if not str(law.get("원본", "")).strip():
                failures.append((mid, name, date, "원문", "원문 비어 있음(상세조회 실패)")); continue
            print(f"  🔍 {mid} {name[:30]}", end=" ", flush=True)
            t1 = time.time()
            try:                                         # ★2026-10-14 법령 단위 격리: 한 건의 예외는 그 건만 실패로 남기고 계속 간다
                ok, new, why = analyze_one(law, old_by_id.get(mid) if args.mock else None)
                if not ok:
                    failures.append((mid, name, date, "분석", why)); print("❌"); continue
                new["_소요초"] = round(time.time() - t1, 1)
                new["_원문글자수"] = len(str(law.get("원본", "")))
                new["_판본출처"] = law.get("판본출처", "")    # 빈칸 = 그날 현행 목록, 값 있음 = 연혁 판본
                new["_프롬프트버전"] = pver                   # ★2026-10-14 캐시 신선도 판정용
                (CACHE / f"result_{mid}.json").write_text(json.dumps(new, ensure_ascii=False), encoding="utf-8")
            except Exception as e:
                failures.append((mid, name, date, "분석", log_exception(mid, name, date, "분석", law, e)))
                print("❌ 예외 → 오류로그 기록, 다음 법령 진행"); continue
            done += 1
            print(f"✅ {new['_소요초']}s")

    print(f"\n· 이번 실행: 신규 {done}건 / 캐시 재사용 {skipped}건 / 실패 {len(failures)}건 / {round(time.time()-t0)}초")
    if ERROR_LOG.exists():
        print(f"  ⚠️ 예외가 기록됨 → {ERROR_LOG.name} (원인 확인용 — 대조 xlsx 와 함께 Claude 에게 올리면 됨)")
    out = build_compare_xlsx(tg, old_by_id, failures, args.mock, model, lvl)
    print(f"✅ {out.name}")


# ─────────────────────────────────────────────────────────────
# 3-2. 승인 추천 (★2026-10-14 신설) — 승인은 사람이 찍고, 여기서는 '추천'과 그 사유만 적는다
# ─────────────────────────────────────────────────────────────
REL_RANK = {"연관높음": 2, "단순관련": 1, "해당없음": 0}


def recommend(old: dict, new: dict, changed: list) -> tuple[str, str]:
    """(추천, 추천사유). 추천 값: 'O 추천' / 'O 추천(소폭)' / '미반영' / '사람 판단'
    규칙(20건 시범의 교훈 — 신뢰도 '높음'인 오답이 있었으므로 종목이 사라진 행은 자동 추천하지 않는다):
      ① 구 관련 종목 중 신에서 사라진 종목이 있으면 → 사람 판단
      ② 우대여부 O→빈칸, 연관도 하향, 신 검토필요 O, 연관높음인데 상세 분석 빈칸 → 사람 판단
      ③ 핵심 판정 열·근거조문이 전부 같으면(해설만 다름) → 미반영
      ④ 핵심 판정이 바뀌었으면(승격·분류·Track 변경) → O 추천 — 근거조문이 비었으면 사람 판단
      ⑤ 그 외(종목 추가·근거조문 변경뿐) → O 추천(소폭)"""
    reasons = []
    old_items, new_items = C.split_items(old.get("관련 종목", "")), C.split_items(new.get("관련 종목", ""))
    lost = sorted(old_items - new_items)
    if lost and "[전종목]" not in new_items:
        n_add = len(new_items - old_items)
        reasons.append(f"구 종목 {len(lost)}개가 신에서 사라짐(+{n_add}개 추가): {', '.join(lost)[:120]}"
                       + (" — 구 분석은 검토필요 O(불확실 표시)였음" if str(old.get("검토필요", "")).strip() == "O" else ""))
    if str(old.get("우대여부", "")).strip() == "O" and str(new.get("우대여부", "")).strip() != "O":
        reasons.append("우대 O→빈칸")
    if REL_RANK.get(str(new.get("연관도", "")).strip(), 1) < REL_RANK.get(str(old.get("연관도", "")).strip(), 1):
        reasons.append(f"연관도 하향 {old.get('연관도', '')}→{new.get('연관도', '')}")
    if str(new.get("검토필요", "")).strip() == "O":
        reasons.append("신 분석이 검토필요 O")
    if str(new.get("연관도", "")).strip() == "연관높음" and not str(new.get("상세 분석 결과", "")).strip():
        reasons.append("연관높음인데 상세 분석 빈칸")
    if reasons:
        return "사람 판단", " / ".join(reasons)
    core_changed = [c for c in C.CORE_COLS if c in changed]
    if not core_changed and "근거조문" not in changed:
        return "미반영", "핵심 판정·종목·근거조문 동일(해설만 다름)"
    if core_changed:
        if not str(new.get("근거조문", "")).strip():
            return "사람 판단", "판정이 바뀌었는데 근거조문 빈칸"
        desc = ", ".join(f"{c} {old.get(c, '') or '(빈칸)'}→{new.get(c, '') or '(빈칸)'}" for c in core_changed if c != "관련 종목")
        added = sorted(new_items - old_items)
        if added:
            desc += f" / 종목 {len(added)}개 추가" if desc else f"종목 {len(added)}개 추가"
        return "O 추천", (desc + f" / 신뢰도 {new.get('AI신뢰도', '')}").strip(" /")
    added = sorted(new_items - old_items)
    bits = []
    if added:
        bits.append(f"종목 {len(added)}개 추가: {', '.join(added)[:100]}")
    if "근거조문" in changed:
        bits.append("근거조문 변경")
    return "O 추천(소폭)", " / ".join(bits) or "경미한 변경"


# ─────────────────────────────────────────────────────────────
# 4. 대조 엑셀
# ─────────────────────────────────────────────────────────────
def build_compare_xlsx(tg, old_by_id, failures, mock, model, lvl) -> Path:
    import pandas as pd
    from openpyxl import load_workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.datavalidation import DataValidation

    recs, trans = [], {c: Counter() for c in ["연관도", "우대여부", "우대분류", "활용도_구분", "검토필요", "AI신뢰도"]}
    fail_ids = {f[0] for f in failures}
    for i, (_, t) in enumerate(tg.iterrows(), 1):
        mid = t["MST_ID"]; old = old_by_id.get(mid, {})
        cp = CACHE / f"result_{mid}.json"
        rec = {"순번": i, "MST_ID": mid, "시행일자": C.norm_date(t["시행일자"]), "법령명": t["법령명"], "점검유형": t.get("점검유형", "")}
        if not cp.exists():
            rec.update({"상태": "실패" if mid in fail_ids else "미실행", "변경필드수": "", "핵심변경": "", "승인": "", "추천": "", "추천사유": "", "검토메모": ""})
            for c in C.ANALYSIS_COLS:
                rec[f"구_{c}"] = old.get(c, ""); rec[f"신_{c}"] = ""
            recs.append(rec); continue
        new = json.loads(cp.read_text(encoding="utf-8"))
        changed = [c for c in C.ANALYSIS_COLS if not C.same_value_col(c, old.get(c, ""), new.get(c, ""))]   # ★2026-10-14 순서 무시
        core_changed = [c for c in C.CORE_COLS if c in changed]
        rcm, why = recommend(old, new, changed)                       # ★2026-10-14 승인 추천
        rec.update({"상태": "성공", "변경필드수": len(changed),
                    "핵심변경": ", ".join(f"{c}: {old.get(c,'')[:12]}→{new.get(c,'')[:12]}" for c in core_changed),
                    "승인": "", "추천": rcm, "추천사유": why, "검토메모": "", "사전밖_제외": new.get("_사전밖_제외", ""),
                    "소요초": new.get("_소요초", ""), "원문글자수": new.get("_원문글자수", ""), "판본출처": new.get("_판본출처", "")})
        for c in C.ANALYSIS_COLS:
            rec[f"구_{c}"] = old.get(c, ""); rec[f"신_{c}"] = new.get(c, "")
        for c in trans:
            trans[c][(old.get(c, "") or "(빈칸)", new.get(c, "") or "(빈칸)")] += 1
        recs.append(rec)

    # 열 순서: 키·상태·승인 → 핵심 열 구/신 쌍 → 나머지 구/신 쌍
    front = ["순번", "MST_ID", "시행일자", "법령명", "점검유형", "상태", "변경필드수", "핵심변경", "승인", "추천", "추천사유", "검토메모", "사전밖_제외", "소요초", "원문글자수", "판본출처"]   # ★2026-10-14 추천·추천사유·판본출처
    pairs = [f"{p}_{c}" for c in C.CORE_COLS for p in ("구", "신")] + \
            [f"{p}_{c}" for c in C.ANALYSIS_COLS if c not in C.CORE_COLS for p in ("구", "신")]
    df = pd.DataFrame(recs).reindex(columns=front + pairs).fillna("")

    ok = df[df["상태"] == "성공"]
    summary = [{"항목": "대상", "값": len(df)}, {"항목": "성공", "값": len(ok)}, {"항목": "실패", "값": len(failures)},
               {"항목": "미실행(캐시 없음)", "값": int((df["상태"] == "미실행").sum())},
               {"항목": "변경 없는 행(성공 중)", "값": int((ok["변경필드수"] == 0).sum()) if len(ok) else 0},
               {"항목": "핵심 판정 바뀐 행", "값": int((ok["핵심변경"] != "").sum()) if len(ok) else 0},
               {"항목": "추천 집계", "값": ", ".join(f"{k} {v}" for k, v in ok["추천"].value_counts().items()) if len(ok) else ""},   # ★2026-10-14
               {"항목": "모델 / thinking_level", "값": f"{model} / {lvl}"}, {"항목": "프롬프트 버전", "값": prompt_version()},
               {"항목": "모드", "값": "★MOCK(테스트)" if mock else "실행"}]
    tr_rows = [{"열": c, "구": a, "신": b, "건수": n} for c, cnt in trans.items() for (a, b), n in sorted(cnt.items(), key=lambda x: -x[1])]
    fail_df = pd.DataFrame(failures, columns=["MST_ID", "법령명", "시행일자", "단계", "사유"])

    path = unique_path(C.WORK / f"재분석_대조_{C.STAMP}{'_MOCK' if mock else ''}.xlsx")   # ★2026-10-14 덮어쓰기·잠금 방지
    with pd.ExcelWriter(path, engine="openpyxl") as w:
        pd.DataFrame(summary).to_excel(w, sheet_name="변경요약", index=False, startrow=0)
        pd.DataFrame(tr_rows).to_excel(w, sheet_name="변경요약", index=False, startrow=len(summary) + 3)
        df.to_excel(w, sheet_name="대조표", index=False)
        fail_df.to_excel(w, sheet_name="실패", index=False)

    # 서식: 헤더, 바뀐 '신_' 칸 노란색, 승인 열 O/X 검증, 틀고정
    wb = load_workbook(path); ws = wb["대조표"]
    hdr = {c.value: c.column for c in ws[1]}
    yellow = PatternFill("solid", fgColor="FFF2CC"); grey = PatternFill("solid", fgColor="F2F2F2")
    for wsx in wb.worksheets:
        for c in wsx[1]:
            c.font = Font(name="맑은 고딕", bold=True, color="FFFFFF", size=10); c.fill = PatternFill("solid", fgColor="1F3864")
        for row in wsx.iter_rows(min_row=2):
            for c in row:
                c.font = Font(name="맑은 고딕", size=9); c.alignment = Alignment(vertical="top", wrap_text=True)
    for r in range(2, ws.max_row + 1):
        if ws.cell(r, hdr["상태"]).value != "성공":
            continue
        for c in C.ANALYSIS_COLS:
            o, n = ws.cell(r, hdr[f"구_{c}"]), ws.cell(r, hdr[f"신_{c}"])
            if not C.same_value_col(c, o.value, n.value):   # ★2026-10-14 순서 무시
                n.fill = yellow
            o.fill = grey if not o.fill.fgColor.rgb or o.fill.fgColor.rgb == "00000000" else o.fill
    dv = DataValidation(type="list", formula1='"O,X"', allow_blank=True); ws.add_data_validation(dv)
    dv.add(f"{get_column_letter(hdr['승인'])}2:{get_column_letter(hdr['승인'])}{max(ws.max_row, 2)}")
    for name, wd in {"순번": 5, "MST_ID": 14, "시행일자": 10, "법령명": 34, "점검유형": 30, "상태": 7, "변경필드수": 7, "핵심변경": 44, "승인": 6, "추천": 11, "추천사유": 46, "검토메모": 24, "사전밖_제외": 18, "소요초": 7, "원문글자수": 9, "판본출처": 22}.items():
        ws.column_dimensions[get_column_letter(hdr[name])].width = wd
    for k, col in hdr.items():
        if k.startswith(("구_", "신_")):
            ws.column_dimensions[get_column_letter(col)].width = 16 if k[2:] in C.CORE_COLS else 22
    ws.freeze_panes = ws.cell(2, hdr["승인"] + 1)
    ws.auto_filter.ref = ws.dimensions
    wb["변경요약"].column_dimensions["A"].width = 24; wb["변경요약"].column_dimensions["B"].width = 30
    for col, wd in zip("ABCDE", [16, 34, 10, 8, 60]): wb["실패"].column_dimensions[col].width = wd
    wb.save(path)
    return path


if __name__ == "__main__":
    main()
